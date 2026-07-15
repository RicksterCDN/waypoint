"""Generate DOCX contract and policy artifacts from Markdown source files."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from .paths import data_path, repo_root


@dataclass(frozen=True)
class GeneratedDocx:
    source_path: Path
    output_path: Path


def generate_docx_documents(
    root: Path | None = None,
    *,
    contracts_output: Path | None = None,
    policies_output: Path | None = None,
) -> list[GeneratedDocx]:
    """Generate DOCX files for contract and policy Markdown sources."""

    root = root or repo_root()
    data_root = data_path(root)
    generated: list[GeneratedDocx] = []

    generated.extend(
        _convert_folder(
            data_root / "contracts" / "source-markdown",
            contracts_output or data_root / "contracts" / "docx",
        )
    )
    generated.extend(
        _convert_folder(
            data_root / "policies" / "source-markdown",
            policies_output or data_root / "policies" / "docx",
        )
    )

    return generated


def _convert_folder(source_folder: Path, output_folder: Path) -> list[GeneratedDocx]:
    if not source_folder.exists():
        raise FileNotFoundError(f"Source folder not found: {source_folder}")

    output_folder.mkdir(parents=True, exist_ok=True)
    generated: list[GeneratedDocx] = []
    for source_path in sorted(source_folder.glob("*.md")):
        output_path = output_folder / f"{source_path.stem}.docx"
        document = _markdown_to_document(source_path.read_text(encoding="utf-8"))
        document.save(output_path)
        generated.append(GeneratedDocx(source_path=source_path, output_path=output_path))
    return generated


def _markdown_to_document(markdown: str) -> Document:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    list_stack: list[int] = []
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            list_stack.clear()
            continue

        if stripped.startswith("# "):
            document.add_heading(stripped[2:].strip(), level=1)
            list_stack.clear()
            continue

        if stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=2)
            list_stack.clear()
            continue

        if stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=3)
            list_stack.clear()
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_runs(paragraph, stripped[2:].strip())
            list_stack.clear()
            continue

        if re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            _add_inline_runs(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            list_stack.clear()
            continue

        paragraph = document.add_paragraph()
        _add_inline_runs(paragraph, stripped)
        list_stack.clear()

    return document


def _add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            _add_text_with_line_breaks(paragraph, part)


def _add_text_with_line_breaks(paragraph, text: str) -> None:
    chunks = text.split("  ")
    for index, chunk in enumerate(chunks):
        if index > 0:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        if chunk:
            paragraph.add_run(chunk)
